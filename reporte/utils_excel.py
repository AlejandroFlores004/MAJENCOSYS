# reporte/utils_excel.py
from io import BytesIO
from typing import List, Tuple, Dict, Any

from openpyxl import Workbook
from django.db.models import F, Sum
from django.db.models.functions import Coalesce

from activity.models import (
    Activity, MemoryMaterial, MemoryManoObra, MemoryHerramienta, MemoryEquipo,
    MemoryRiesgos, MemoryCalidad, MemoryAmbiental, MemoryHidrologica
)

# Import del modelo Schedule, sea CamelCase o minúscula
try:
    from schedule.models import Schedule  # Clase típica en CamelCase
except Exception:  # pragma: no cover - fallback defensivo
    from schedule.models import schedule as Schedule  # Modelo declarado en minúscula


# ------------------------------- helpers comunes ----------------------------

def _wb_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _totales_costos(activity_ids: List[int]) -> Dict[int, float]:
    """
    Suma de todos los costos por actividad (para la columna 'total').
    Retorna {activity_id: total_usd}
    """
    if not activity_ids:
        return {}

    def agg(qs):
        return {x["activity_id"]: float(x["total"] or 0) for x in qs}

    mat = (MemoryMaterial.objects.filter(activity_id__in=activity_ids)
           .annotate(sub=F("quantity") * F("material__price"))
           .values("activity_id").annotate(total=Coalesce(Sum("sub"), 0)))
    mo = (MemoryManoObra.objects.filter(activity_id__in=activity_ids)
          .annotate(sub=F("manoobra__jornada") * (1 + F("prestaciones")) * F("rendimiento"))
          .values("activity_id").annotate(total=Coalesce(Sum("sub"), 0)))
    hrr = (MemoryHerramienta.objects.filter(activity_id__in=activity_ids)
           .annotate(sub=F("herramienta__costodia") * F("rendimiento"))
           .values("activity_id").annotate(total=Coalesce(Sum("sub"), 0)))
    eqp = (MemoryEquipo.objects.filter(activity_id__in=activity_ids)
           .annotate(sub=F("equipo__costodia") * F("rendimiento"))
           .values("activity_id").annotate(total=Coalesce(Sum("sub"), 0)))
    rsg = (MemoryRiesgos.objects.filter(activity_id__in=activity_ids)
           .values("activity_id").annotate(total=Coalesce(Sum("costo"), 0)))
    cld = (MemoryCalidad.objects.filter(activity_id__in=activity_ids)
           .annotate(sub=F("calidad__precio") * F("cantidad"))
           .values("activity_id").annotate(total=Coalesce(Sum("sub"), 0)))
    amb = (MemoryAmbiental.objects.filter(activity_id__in=activity_ids)
           .values("activity_id").annotate(total=Coalesce(Sum("valor"), 0)))
    hdg = (MemoryHidrologica.objects.filter(activity_id__in=activity_ids)
           .values("activity_id").annotate(total=Coalesce(Sum("costo"), 0)))

    res: Dict[int, float] = {}
    for m in [agg(mat), agg(mo), agg(hrr), agg(eqp), agg(rsg), agg(cld), agg(amb), agg(hdg)]:
        for k, v in m.items():
            res[k] = res.get(k, 0.0) + v
    return res


def _duration(start, end):
    if start and end:
        return (end - start).days
    return None


# ----------------------------- motor de reportes ----------------------------

DEFAULT_COLUMNS = ["actividad", "estado", "inicio", "fin"]
ALL_COLUMNS = ["actividad", "estado", "inicio", "fin", "duracion", "total"]


def build_custom_report(project, payload: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    """
    Devuelve (headers, rows) según 'payload':
      payload = {
        "columns": ["actividad","estado","inicio","fin","duracion","total"],
        "filters": {
          "estado": ["P","E","F"],
          "ini_desde": "YYYY-MM-DD" | null,
          "ini_hasta": "YYYY-MM-DD" | null,
          "fin_desde":  "YYYY-MM-DD" | null,
          "fin_hasta":  "YYYY-MM-DD" | null
        },
        "order": "actividad|inicio|fin|-fin|total|-total",
        "group_by": ""|"estado"   # (de momento informativo)
      }
    """
    cols = payload.get("columns") or DEFAULT_COLUMNS
    cols = [c for c in cols if c in ALL_COLUMNS]

    # Base: cronogramas del proyecto con su actividad
    sch = Schedule.objects.filter(activity__project=project).select_related("activity")

    # Filtros
    flt = (payload.get("filters") or {})
    estados = flt.get("estado") or []
    if estados:
        sch = sch.filter(status__in=estados)
    if flt.get("ini_desde"):
        sch = sch.filter(start_date__gte=flt["ini_desde"])
    if flt.get("ini_hasta"):
        sch = sch.filter(start_date__lte=flt["ini_hasta"])
    if flt.get("fin_desde"):
        sch = sch.filter(end_date__gte=flt["fin_desde"])
    if flt.get("fin_hasta"):
        sch = sch.filter(end_date__lte=flt["fin_hasta"])

    # Orden
    order = payload.get("order") or ""
    if order in {"actividad", "inicio", "fin", "-fin"}:
        map_order = {
            "actividad": "activity__name",
            "inicio": "start_date",
            "fin": "end_date",
            "-fin": "-end_date",
        }
        sch = sch.order_by(map_order[order])
        sch_iter = list(sch)
    elif order in {"total", "-total"}:
        ids_for_totals = list(sch.values_list("activity_id", flat=True))
        totals_map = _totales_costos(ids_for_totals)
        sch_iter = sorted(list(sch), key=lambda s: totals_map.get(s.activity_id, 0.0), reverse=(order == "-total"))
    else:
        sch = sch.order_by("activity__name")
        sch_iter = list(sch)

    # Precalcular totales si se necesitan
    totals: Dict[int, float] = {}
    if "total" in cols:
        ids = [s.activity_id for s in sch_iter]
        totals = _totales_costos(ids)

    # Cabeceras
    headers_map = {
        "actividad": "Actividad",
        "estado": "Estado",
        "inicio": "Inicio",
        "fin": "Fin",
        "duracion": "Duración (días)",
        "total": "Total (USD)",
    }
    headers = [headers_map[c] for c in cols]

    # Filas
    rows: List[List[Any]] = []
    for s in sch_iter:
        a = s.activity
        row = []
        for c in cols:
            if c == "actividad":
                row.append(a.name)
            elif c == "estado":
                row.append(s.get_status_display())
            elif c == "inicio":
                row.append(s.start_date)
            elif c == "fin":
                row.append(s.end_date)
            elif c == "duracion":
                row.append(_duration(s.start_date, s.end_date))
            elif c == "total":
                row.append(round(totals.get(a.id, 0.0), 2))
        rows.append(row)

    return headers, rows


# ------------------------- generadores de Excel existentes ------------------

def excel_libro_mayor(project) -> bytes:
    payload = {
        "columns": ["actividad", "estado", "inicio", "fin", "total"],
        "filters": {},
        "order": "actividad",
    }
    headers, rows = build_custom_report(project, payload)

    wb = Workbook()
    ws = wb.active
    ws.title = "Libro mayor"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    return _wb_bytes(wb)


def excel_actividades_detalle(project) -> bytes:
    payload = {
        "columns": ["actividad", "estado", "inicio", "fin", "duracion"],
        "filters": {},
        "order": "actividad",
    }
    headers, rows = build_custom_report(project, payload)

    wb = Workbook()
    ws = wb.active
    ws.title = "Actividades"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    return _wb_bytes(wb)


def excel_actividades_por_estado(project, estado: str) -> bytes:
    payload = {
        "columns": ["actividad", "estado", "inicio", "fin"],
        "filters": {"estado": [estado]},
        "order": "actividad",
    }
    headers, rows = build_custom_report(project, payload)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Estado {estado}"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    return _wb_bytes(wb)


def excel_actividades_por_fecha(project, fecha: str, tipo: str = "fin") -> bytes:
    if tipo == "inicio":
        flt = {"ini_desde": fecha, "ini_hasta": fecha}
        cols = ["actividad", "estado", "inicio", "duracion"]
    else:
        flt = {"fin_desde": fecha, "fin_hasta": fecha}
        cols = ["actividad", "estado", "fin", "duracion"]

    payload = {"columns": cols, "filters": flt, "order": "actividad"}
    headers, rows = build_custom_report(project, payload)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Por fecha {tipo}"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    return _wb_bytes(wb)


def excel_from_payload(project, payload: Dict[str, Any]) -> bytes:
    """Genera Excel desde el payload del constructor o plantilla guardada."""
    headers, rows = build_custom_report(project, payload)
    wb = Workbook()
    ws = wb.active
    ws.title = payload.get("name") or "Reporte"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    return _wb_bytes(wb)
