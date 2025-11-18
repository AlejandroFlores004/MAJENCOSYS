from decimal import Decimal
from collections import defaultdict
from io import BytesIO

from django.db.models import Sum, F
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
    Border,
    Side,
    NamedStyle,
)
from openpyxl.utils import get_column_letter

from activity.models import Activity
from schedule.models import Schedule  # <- tu modelo de cronograma

from activity.models import (
    MemoryMaterial,
    MemoryEquipo,
    MemoryHerramienta,
    MemoryHidrologica,
    MemoryAmbiental,
    # si más adelante quieres mano de obra, aquí la importas
)

# ---------------------------------------------------------------------
# helpers de estilo
# ---------------------------------------------------------------------


def _money_style():
    s = NamedStyle(name="MoneyUSD")
    s.number_format = u'[$$-409] #,##0.00'
    s.alignment = Alignment(vertical="center")
    return s


def _date_style():
    s = NamedStyle(name="DateISO")
    s.number_format = "yyyy-mm-dd"
    s.alignment = Alignment(vertical="center")
    return s


def _header_style(cell):
    cell.font = Font(bold=True, color="000000")
    cell.fill = PatternFill("solid", fgColor="E6EAF2")
    thin = Side(style="thin", color="DDDDDD")
    cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def _body_border(cell):
    thin = Side(style="thin", color="DDDDDD")
    cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)

def _autowidth(ws):
    """
    Ajusta el ancho de las columnas según el contenido.
    Para columnas de fecha (Inicio, Fin, Fecha...) les da un ancho mínimo mayor
    para evitar los #### en Excel.
    """
    widths = defaultdict(int)

    # medir contenido
    for row in ws.iter_rows(values_only=True):
        for i, v in enumerate(row, start=1):
            if v is None:
                continue
            s = str(v)
            l = len(s) + (2 if any(c in s for c in (".", ",")) else 0)
            widths[i] = max(widths[i], min(l, 50))

    # cabecera para detectar columnas de fecha
    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), [])
    for idx, head in enumerate(headers, start=1):
        if not head:
            continue
        h = str(head).lower()
        if "inicio" in h or "fin" in h or "fecha" in h:
            # aseguro buen ancho para fechas
            widths[idx] = max(widths.get(idx, 0), 14)

    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(10, w or 10)


# ---------------------------------------------------------------------
# columnas permitidas
# ---------------------------------------------------------------------

_ALLOWED_COLUMNS = {
    "actividad": "Actividad",
    "estado": "Estado",
    "inicio": "Inicio",
    "fin": "Fin",
    "duracion": "Duración",
    "total": "Total (USD)",
}

_DEFAULT_COLUMNS = ["actividad", "estado", "inicio", "fin", "duracion", "total"]


def _normalize_payload(payload: dict) -> dict:
    payload = payload or {}
    cols = payload.get("columns") or _DEFAULT_COLUMNS
    cols = [c for c in cols if c in _ALLOWED_COLUMNS]
    if not cols:
        cols = _DEFAULT_COLUMNS

    order = payload.get("order") or "actividad"
    if order not in _ALLOWED_COLUMNS:
        order = "actividad"

    filters = payload.get("filters") or {}

    name = (payload.get("name") or "Custom").strip() or "Custom"
    group_by = payload.get("group_by") or ""

    return {
        "columns": cols,
        "order": order,
        "filters": filters,
        "name": name,
        "group_by": group_by,
    }


# ---------------------------------------------------------------------
# obtención de datos base
# ---------------------------------------------------------------------


def _activity_queryset(project):
    # trae las actividades del proyecto
    return Activity.objects.filter(project=project).select_related("project")


def _schedule_map(activity_ids):
    """
    Devuelve un dict {activity_id: (start_date, end_date, status)} tomando el ÚLTIMO
    schedule de cada actividad (por id).
    """
    out = {}
    for sch in (
        Schedule.objects.filter(activity_id__in=activity_ids)
        .order_by("activity_id", "-id")
        .select_related("activity")
    ):
        # solo guardamos el primero que veamos por actividad (el más reciente por -id)
        if sch.activity_id not in out:
            out[sch.activity_id] = (sch.start_date, sch.end_date, sch.status)
    return out


def _money_or_zero(x):
    try:
        return Decimal(x or 0)
    except Exception:
        return Decimal(0)


def _totales_por_actividad(project, activity_ids):
    """
    Suma los diferentes tipos de memoria por actividad y devuelve
    {activity_id: Decimal(totalUSD)}
    """
    # materiales
    mats = (
        MemoryMaterial.objects.filter(activity_id__in=activity_ids)
        .annotate(sub=F("quantity") * F("material__price"))
        .values("activity_id")
        .annotate(total=Sum("sub"))
    )
    m_map = {r["activity_id"]: _money_or_zero(r["total"]) for r in mats}

    # equipo
    eqs = (
        MemoryEquipo.objects.filter(activity_id__in=activity_ids)
        .annotate(sub=F("rendimiento") * F("equipo__costodia"))
        .values("activity_id")
        .annotate(total=Sum("sub"))
    )
    e_map = {r["activity_id"]: _money_or_zero(r["total"]) for r in eqs}

    # herramienta
    hrs = (
        MemoryHerramienta.objects.filter(activity_id__in=activity_ids)
        .annotate(sub=F("rendimiento") * F("herramienta__costodia"))
        .values("activity_id")
        .annotate(total=Sum("sub"))
    )
    h_map = {r["activity_id"]: _money_or_zero(r["total"]) for r in hrs}

    # hidrológica
    hids = (
        MemoryHidrologica.objects.filter(activity_id__in=activity_ids)
        .values("activity_id")
        .annotate(total=Sum("costo"))
    )
    hd_map = {r["activity_id"]: _money_or_zero(r["total"]) for r in hids}

    # ambiental
    ambs = (
        MemoryAmbiental.objects.filter(activity_id__in=activity_ids)
        .values("activity_id")
        .annotate(total=Sum("valor"))
    )
    a_map = {r["activity_id"]: _money_or_zero(r["total"]) for r in ambs}

    out = {}
    for aid in activity_ids:
        out[aid] = (
            m_map.get(aid, 0)
            + e_map.get(aid, 0)
            + h_map.get(aid, 0)
            + hd_map.get(aid, 0)
            + a_map.get(aid, 0)
        )
    return out


# ---------------------------------------------------------------------
# API principal usada por las views
# ---------------------------------------------------------------------


def build_custom_report(project, payload):
    """
    Devuelve (headers, rows)
    rows es lista de listas con el mismo orden de headers.
    """
    cfg = _normalize_payload(payload)
    qs = _activity_queryset(project)

    filters = cfg["filters"]

    # ---- filtro por estado (viene del schedule, no de activity) ----
    estados = filters.get("estado")
    if estados:
        qs = qs.filter(schedules__status__in=estados).distinct()


    # ---- filtro por fechas usando schedule ----
    # Compatibilidad: soportamos
    #  - fecha_desde / fecha_hasta  (rango aplicado de forma global)
    #  - ini_desde / ini_hasta      (rango sobre start_date)
    #  - fin_desde / fin_hasta      (rango sobre end_date)
    fecha_desde = filters.get("fecha_desde")
    fecha_hasta = filters.get("fecha_hasta")

    ini_desde = filters.get("ini_desde")
    ini_hasta = filters.get("ini_hasta")
    fin_desde = filters.get("fin_desde")
    fin_hasta = filters.get("fin_hasta")

    if fecha_desde:
        qs = qs.filter(schedules__start_date__gte=fecha_desde)
    if fecha_hasta:
        qs = qs.filter(schedules__end_date__lte=fecha_hasta)

    if ini_desde:
        qs = qs.filter(schedules__start_date__gte=ini_desde)
    if ini_hasta:
        qs = qs.filter(schedules__start_date__lte=ini_hasta)

    if fin_desde:
        qs = qs.filter(schedules__end_date__gte=fin_desde)
    if fin_hasta:
        qs = qs.filter(schedules__end_date__lte=fin_hasta)

    if any([fecha_desde, fecha_hasta, ini_desde, ini_hasta, fin_desde, fin_hasta]):
        qs = qs.distinct()

###
    # orden básico por nombre
    order = cfg["order"]
    if order == "actividad":
        qs = qs.order_by("name")
    # si es por inicio / fin / total lo haremos en memoria

    act_list = list(qs)
    act_ids = [a.id for a in act_list]

    sch_map = _schedule_map(act_ids)
    totals_map = _totales_por_actividad(project, act_ids)

    cols = cfg["columns"]
    headers = [_ALLOWED_COLUMNS[c] for c in cols]

    rows_raw = []
    for a in act_list:
        sch_info = sch_map.get(a.id)
        if sch_info:
            inicio, fin, est = sch_info
        else:
            inicio, fin, est = (None, None, "")

        dur = None
        if inicio and fin:
            dur = (fin - inicio).days

        total = totals_map.get(a.id, Decimal(0))

        values = {
            "actividad": a.name,
            "estado": est,
            "inicio": inicio,
            "fin": fin,
            "duracion": dur,
            "total": total,
        }
        rows_raw.append([values[c] for c in cols])

    # ordenamientos que necesitan datos ya construidos
    if order in ("inicio", "fin", "total"):
        if order in cols:
            idx = cols.index(order)
            rows_raw.sort(key=lambda r: (r[idx] is None, r[idx]))

    return headers, rows_raw


def excel_from_payload(project, payload):
    headers, rows = build_custom_report(project, payload)

    wb = Workbook()
    ws = wb.active
    ws.title = (payload or {}).get("name") or "Custom"

    money = _money_style()
    date = _date_style()
    # registramos estilos (si ya existen no nos caemos)
    try:
        wb.add_named_style(money)
    except ValueError:
        pass
    try:
        wb.add_named_style(date)
    except ValueError:
        pass

    # encabezados
    ws.append(headers)
    for c in ws[1]:
        _header_style(c)

    # filas
    for row in rows:
        ws.append(row)

    # aplicar estilos por nombre de columna
    col_index = {headers[i].lower(): i + 1 for i in range(len(headers))}

    if "inicio" in col_index:
        col = col_index["inicio"]
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.style = "DateISO"
                _body_border(c)

    if "fin" in col_index:
        col = col_index["fin"]
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.style = "DateISO"
                _body_border(c)

    if "total (usd)" in col_index:
        col = col_index["total (usd)"]
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2):
            for c in cell:
                c.style = "MoneyUSD"
                _body_border(c)

    # bordes al resto
    max_row = ws.max_row
    max_col = ws.max_column
    for r in ws.iter_rows(min_row=2, max_row=max_row, min_col=1, max_col=max_col):
        for c in r:
            if not c.has_style:
                _body_border(c)

    # fila TOTAL si hay columna de total
    if "total (usd)" in col_index and ws.max_row >= 2:
        total_col = col_index["total (usd)"]
        ws.append(
            [None] * (total_col - 2)
            + [
                "TOTAL:",
                f"=SUM({get_column_letter(total_col)}2:{get_column_letter(total_col)}{ws.max_row})",
            ]
        )
        last_row = ws.max_row
        ws.cell(row=last_row, column=total_col).style = "MoneyUSD"
        ws.cell(row=last_row, column=total_col - 1).font = Font(bold=True)

    # filtro y freeze
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    _autowidth(ws)

    out = BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


# ---------------------------------------------------------------------
# funciones de compatibilidad para las vistas "rápidas"
# ---------------------------------------------------------------------


def excel_libro_mayor(project):
    payload = {
        "name": "Libro mayor",
        "columns": ["actividad", "estado", "inicio", "fin", "duracion", "total"],
        "order": "actividad",
    }
    return excel_from_payload(project, payload)


def excel_actividades_detalle(project):
    payload = {
        "name": "Actividades detalle",
        "columns": ["actividad", "estado", "inicio", "fin", "duracion", "total"],
        "order": "inicio",
    }
    return excel_from_payload(project, payload)


def excel_actividades_por_fecha(project, fecha, tipo="fin"):
    # tipo: "inicio" o "fin"
    if tipo not in ("inicio", "fin"):
        tipo = "fin"
    payload = {
        "name": f"Actividades por fecha {tipo} {fecha}",
        "columns": ["actividad", "estado", "inicio", "fin", "duracion", "total"],
        "filters": {
            "fecha_desde": fecha,
            "fecha_hasta": fecha,
        },
        "order": tipo,
    }
    return excel_from_payload(project, payload)
